"""
app/importers/database_writer.py

MotorTracker
Inserta registros nuevos en MASTER TRACK MOTORES.xlsx
conservando la estructura y el formato del archivo.
"""

from __future__ import annotations

import shutil
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from app.core.exceptions import (
    BackupError,
    DatabaseNotFoundError,
    DatabaseWriteError,
    MissingColumnError,
)
from app.models.motor_record import MotorRecord
from app.utils.logger import LoggerManager


class DatabaseWriter:
    """
    Escribe registros MotorRecord en la base de datos Excel.

    La clase conserva:

    - Estilos
    - Bordes
    - Colores
    - Fuentes
    - Formatos numéricos
    - Validaciones
    - Fórmulas existentes
    - Archivos con macros, cuando se utiliza .xlsm
    """

    HEADER_ALIASES = {
        "MODEL": "MODEL",
        "PN": "PN",
        "PART NUMBER": "PN",
        "SUPPLIER": "SUPPLIER",
        "TRACEABILITY": "TRACEABILITY",
        "TRACEBILITY": "TRACEABILITY",
        "NO. MOTORS": "NO. MOTORS",
        "NO MOTORS": "NO. MOTORS",
        "NUMBER OF MOTORS": "NO. MOTORS",
        "VIN": "VIN",
        "NCT": "NCT",
        "N/C": "NCT",
        "N/C #": "NCT",
        "SDA/RMA": "SDA/RMA",
        "SDA / RMA": "SDA/RMA",
        "SDA/ RMA": "SDA/RMA",
        "SDA RMA": "SDA/RMA",
        "DTC": "DTC",
        "ISSUE": "ISSUE",
        "LOCATION": "LOCATION",
        "COMMENTS": "COMMENTS",
        "DISPOSITION": "DISPOSITION",
        "ICA": "ICA",
        "RCA": "RCA",
        "PCA": "PCA",
        "COMMENTS2": "COMMENTS2",
        "COMMENTS 2": "COMMENTS2",
    }

    ESSENTIAL_COLUMNS = {
        "MODEL",
        "PN",
        "SUPPLIER",
        "TRACEABILITY",
        "NO. MOTORS",
        "VIN",
        "NCT",
        "ISSUE",
    }

    def __init__(
        self,
        database_path: Path | str,
        sheet_name: str | None = None,
        backup_folder: Path | str = "Backups",
    ) -> None:
        self.database_path = Path(database_path)
        self.sheet_name = sheet_name
        self.backup_folder = Path(backup_folder)

        self.header_row: int = 0
        self.column_map: dict[str, int] = {}
        self.last_written_rows: list[int] = []

    # =========================================================
    # PUBLIC
    # =========================================================

    def write(
        self,
        records: Iterable[MotorRecord],
        create_backup: bool = True,
    ) -> list[int]:
        """
        Inserta una colección de registros en la base de datos.

        Args:
            records:
                Registros que serán insertados.

            create_backup:
                Cuando es True, crea una copia de seguridad antes
                de modificar la base.

        Returns:
            Lista con los números de fila insertados.
        """

        record_list = list(records)

        if not record_list:
            LoggerManager.info(
                "No hay registros nuevos para escribir."
            )
            return []

        self._validate_database()

        backup_path: Path | None = None

        try:
            if create_backup:
                backup_path = self.create_backup()

                LoggerManager.info(
                    f"Backup creado: {backup_path.name}"
                )

            keep_vba = (
                self.database_path.suffix.lower() == ".xlsm"
            )

            workbook = load_workbook(
                filename=self.database_path,
                keep_vba=keep_vba,
                data_only=False,
                keep_links=True,
            )

            worksheet = self._select_worksheet(workbook.worksheets)

            self.header_row = self._find_header_row(worksheet)
            self.column_map = self._build_column_map(
                worksheet,
                self.header_row,
            )

            self._validate_columns()

            first_available_row = self._find_next_available_row(
                worksheet
            )

            inserted_rows: list[int] = []

            for index, record in enumerate(record_list):
                row_number = first_available_row + index

                self._prepare_new_row(
                    worksheet=worksheet,
                    row_number=row_number,
                )

                self._write_record(
                    worksheet=worksheet,
                    row_number=row_number,
                    record=record,
                )

                inserted_rows.append(row_number)

            workbook.save(self.database_path)
            workbook.close()

            self.last_written_rows = inserted_rows

            LoggerManager.info(
                f"Registros insertados: {len(inserted_rows)}"
            )

            LoggerManager.info(
                f"Base actualizada: {self.database_path.name}"
            )

            return inserted_rows

        except (
            DatabaseNotFoundError,
            MissingColumnError,
            DatabaseWriteError,
            BackupError,
        ):
            raise

        except PermissionError as error:
            LoggerManager.exception(
                "No se pudo escribir en la base."
            )

            raise DatabaseWriteError(
                "No fue posible guardar la base de datos. "
                "Verifica que el archivo no esté abierto en Excel "
                "y que tengas permisos de escritura."
            ) from error

        except Exception as error:
            LoggerManager.exception(
                "Ocurrió un error escribiendo en la base."
            )

            raise DatabaseWriteError(
                f"No fue posible actualizar la base: {error}"
            ) from error

    def write_record(
        self,
        record: MotorRecord,
        create_backup: bool = True,
    ) -> int:
        """
        Inserta un solo registro y devuelve su número de fila.
        """

        rows = self.write(
            records=[record],
            create_backup=create_backup,
        )

        if not rows:
            raise DatabaseWriteError(
                "El registro no pudo ser insertado."
            )

        return rows[0]

    def create_backup(self) -> Path:
        """
        Crea una copia de seguridad con fecha y hora.
        """

        self._validate_database()

        try:
            self.backup_folder.mkdir(
                parents=True,
                exist_ok=True,
            )

            timestamp = datetime.now().strftime(
                "%Y%m%d_%H%M%S_%f"
            )

            backup_name = (
                f"{self.database_path.stem}_"
                f"{timestamp}"
                f"{self.database_path.suffix}"
            )

            backup_path = self.backup_folder / backup_name

            shutil.copy2(
                self.database_path,
                backup_path,
            )

            return backup_path

        except Exception as error:
            raise BackupError(
                f"No fue posible crear el backup: {error}"
            ) from error

    def get_last_written_rows(self) -> list[int]:
        """
        Devuelve las filas insertadas en la última operación.
        """

        return list(self.last_written_rows)

    def summary(self) -> dict[str, Any]:
        """
        Devuelve información de la última operación.
        """

        return {
            "database_path": str(self.database_path),
            "sheet_name": self.sheet_name or "",
            "header_row": self.header_row,
            "column_map": dict(self.column_map),
            "written_rows": list(self.last_written_rows),
            "written_records": len(self.last_written_rows),
        }

    # =========================================================
    # DATABASE VALIDATION
    # =========================================================

    def _validate_database(self) -> None:
        """
        Comprueba que la base exista y tenga una extensión válida.
        """

        if not self.database_path.exists():
            raise DatabaseNotFoundError(
                f"No se encontró la base de datos: "
                f"{self.database_path}"
            )

        if not self.database_path.is_file():
            raise DatabaseWriteError(
                "La ruta de la base no corresponde a un archivo."
            )

        if self.database_path.suffix.lower() not in {
            ".xlsx",
            ".xlsm",
        }:
            raise DatabaseWriteError(
                "La base debe tener extensión .xlsx o .xlsm."
            )

    # =========================================================
    # WORKSHEET
    # =========================================================

    def _select_worksheet(
        self,
        worksheets: list[Worksheet],
    ) -> Worksheet:
        """
        Selecciona la hoja indicada o detecta automáticamente
        la que contiene la base de motores.
        """

        if not worksheets:
            raise DatabaseWriteError(
                "La base de datos no contiene hojas."
            )

        if self.sheet_name:
            for worksheet in worksheets:
                if (
                    worksheet.title.strip().casefold()
                    == self.sheet_name.strip().casefold()
                ):
                    return worksheet

            raise DatabaseWriteError(
                f"No se encontró la hoja '{self.sheet_name}'."
            )

        best_worksheet: Worksheet | None = None
        best_score = 0

        for worksheet in worksheets:
            score = self._calculate_sheet_score(worksheet)

            if score > best_score:
                best_worksheet = worksheet
                best_score = score

        if best_worksheet is None or best_score < 5:
            raise DatabaseWriteError(
                "No se encontró una hoja con la estructura "
                "esperada de MASTER TRACK MOTORES."
            )

        self.sheet_name = best_worksheet.title

        return best_worksheet

    def _calculate_sheet_score(
        self,
        worksheet: Worksheet,
    ) -> int:
        """
        Calcula cuántos encabezados conocidos contiene una hoja.
        """

        highest_score = 0
        max_rows = min(worksheet.max_row, 30)

        for row_number in range(1, max_rows + 1):
            detected_headers: set[str] = set()

            for cell in worksheet[row_number]:
                normalized = self._normalize_header(cell.value)
                canonical = self.HEADER_ALIASES.get(normalized)

                if canonical:
                    detected_headers.add(canonical)

            highest_score = max(
                highest_score,
                len(detected_headers),
            )

        return highest_score

    def _find_header_row(
        self,
        worksheet: Worksheet,
    ) -> int:
        """
        Localiza la fila que contiene los encabezados.
        """

        best_row = 0
        best_score = 0
        max_rows = min(worksheet.max_row, 30)

        for row_number in range(1, max_rows + 1):
            detected_headers: set[str] = set()

            for cell in worksheet[row_number]:
                normalized = self._normalize_header(cell.value)
                canonical = self.HEADER_ALIASES.get(normalized)

                if canonical:
                    detected_headers.add(canonical)

            score = len(detected_headers)

            if score > best_score:
                best_row = row_number
                best_score = score

        if best_row == 0 or best_score < 5:
            raise MissingColumnError(
                "No se encontró la fila de encabezados "
                "en la base de datos."
            )

        return best_row

    def _build_column_map(
        self,
        worksheet: Worksheet,
        header_row: int,
    ) -> dict[str, int]:
        """
        Relaciona encabezados con números de columna.
        """

        column_map: dict[str, int] = {}

        for cell in worksheet[header_row]:
            normalized = self._normalize_header(cell.value)
            canonical = self.HEADER_ALIASES.get(normalized)

            if canonical and canonical not in column_map:
                column_map[canonical] = cell.column

        return column_map

    def _validate_columns(self) -> None:
        """
        Verifica que existan las columnas mínimas necesarias.
        """

        missing_columns = sorted(
            self.ESSENTIAL_COLUMNS - set(self.column_map)
        )

        if missing_columns:
            raise MissingColumnError(
                "Faltan columnas obligatorias en la base: "
                + ", ".join(missing_columns)
            )

    # =========================================================
    # ROW MANAGEMENT
    # =========================================================

    def _find_next_available_row(
        self,
        worksheet: Worksheet,
    ) -> int:
        """
        Encuentra la primera fila disponible debajo de los datos.
        """

        important_columns = (
            "NCT",
            "VIN",
            "TRACEABILITY",
            "PN",
            "ISSUE",
        )

        for row_number in range(
            worksheet.max_row,
            self.header_row,
            -1,
        ):
            for column_name in important_columns:
                column_number = self.column_map.get(column_name)

                if not column_number:
                    continue

                value = worksheet.cell(
                    row=row_number,
                    column=column_number,
                ).value

                if self._as_text(value):
                    return row_number + 1

        return self.header_row + 1

    def _prepare_new_row(
        self,
        worksheet: Worksheet,
        row_number: int,
    ) -> None:
        """
        Copia el formato de la fila anterior a la nueva fila.
        """

        template_row = row_number - 1

        if template_row <= self.header_row:
            return

        worksheet.row_dimensions[row_number].height = (
            worksheet.row_dimensions[template_row].height
        )

        for column_number in range(
            1,
            worksheet.max_column + 1,
        ):
            source_cell = worksheet.cell(
                row=template_row,
                column=column_number,
            )

            target_cell = worksheet.cell(
                row=row_number,
                column=column_number,
            )

            self._copy_cell_style(
                source=source_cell,
                target=target_cell,
            )

    @staticmethod
    def _copy_cell_style(
        source: Cell,
        target: Cell,
    ) -> None:
        """
        Copia solamente el formato de una celda.
        No copia el valor para evitar duplicar información.
        """

        if source.has_style:
            target._style = copy(source._style)

        if source.number_format:
            target.number_format = source.number_format

        if source.font:
            target.font = copy(source.font)

        if source.fill:
            target.fill = copy(source.fill)

        if source.border:
            target.border = copy(source.border)

        if source.alignment:
            target.alignment = copy(source.alignment)

        if source.protection:
            target.protection = copy(source.protection)

    # =========================================================
    # RECORD WRITING
    # =========================================================

    def _write_record(
        self,
        worksheet: Worksheet,
        row_number: int,
        record: MotorRecord,
    ) -> None:
        """
        Escribe los campos de un MotorRecord en una fila.
        """

        values = {
            "MODEL": record.model,
            "PN": record.part_number,
            "SUPPLIER": record.supplier,
            "TRACEABILITY": record.traceability,
            "NO. MOTORS": record.no_motors,
            "VIN": record.vin,
            "NCT": record.nct,
            "SDA/RMA": self._combine_sda_rma(record),
            "DTC": record.dtc,
            "ISSUE": record.issue,
            "LOCATION": record.location,
            "COMMENTS": record.comments,
            "DISPOSITION": record.disposition,
            "ICA": record.ica,
            "RCA": record.rca,
            "PCA": record.pca,
            "COMMENTS2": record.comments2,
        }

        for column_name, value in values.items():
            column_number = self.column_map.get(column_name)

            if not column_number:
                continue

            worksheet.cell(
                row=row_number,
                column=column_number,
            ).value = self._clean_excel_value(value)

    @staticmethod
    def _combine_sda_rma(
        record: MotorRecord,
    ) -> str:
        """
        Combina SDA y RMA cuando ambos contienen información.
        """

        sda = DatabaseWriter._as_text(record.sda)
        rma = DatabaseWriter._as_text(record.rma)

        if sda and rma:
            if sda.casefold() == rma.casefold():
                return sda

            return f"SDA: {sda} / RMA: {rma}"

        return sda or rma

    # =========================================================
    # HELPERS
    # =========================================================

    @staticmethod
    def _normalize_header(value: Any) -> str:
        """
        Normaliza encabezados de Excel.
        """

        if value is None:
            return ""

        text = str(value)
        text = text.replace("\n", " ")
        text = text.replace("\r", " ")
        text = text.replace("\xa0", " ")
        text = " ".join(text.split())

        return text.strip().upper()

    @staticmethod
    def _as_text(value: Any) -> str:
        """
        Convierte un valor a texto limpio.
        """

        if value is None:
            return ""

        if isinstance(value, float) and value.is_integer():
            return str(int(value))

        text = str(value)
        text = text.replace("\xa0", " ")

        return text.strip()

    @staticmethod
    def _clean_excel_value(value: Any) -> Any:
        """
        Limpia valores antes de escribirlos en Excel.
        """

        if value is None:
            return ""

        if isinstance(value, str):
            value = value.replace("\x00", "")
            value = value.replace("\xa0", " ")
            return value.strip()

        return value